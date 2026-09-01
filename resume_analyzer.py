import io
import os
from typing import List, Optional
from pydantic import BaseModel, Field
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document

# Define schemas for structured JSON output from Gemini
class CategoryScore(BaseModel):
    category: str = Field(description="The name of the category evaluated, e.g. Skills Alignment, Experience Relevance, Formatting & Structure, Education")
    score: int = Field(description="Score for this category from 0 to 100")
    feedback: str = Field(description="Specific visual summary feedback for this score")

class KeywordMatch(BaseModel):
    keyword: str = Field(description="Important keyword or skill required by the JD")
    present: bool = Field(description="True if the keyword is found in the resume, False if missing")
    importance: str = Field(description="The level of importance of this keyword: high, medium, or low")

class BulletPointImprovement(BaseModel):
    original: str = Field(description="Original bullet point or experience description from the resume")
    improved: str = Field(description="A stronger rewritten version of the bullet point using strong action verbs. Do NOT invent metrics, percentages, or experience. Only use facts from the resume.")
    reason: str = Field(description="Identify why the original is weak and explain the improvements made. If there is not enough information to quantify, explicitly state what measurable/quantifiable information the user should add themselves.")

class ResumeAnalysis(BaseModel):
    match_score: int = Field(description="Overall ATS match score between the resume and the job description, from 0 to 100")
    rating: str = Field(description="Overall fit description, e.g., 'Excellent Match', 'Strong Match', 'Partial Match', 'Low Match'")
    summary: str = Field(description="A high-level professional summary of the resume's match against the job description (2-3 sentences)")
    category_scores: List[CategoryScore] = Field(description="Detailed scoring for different dimensions of the resume")
    strengths: List[str] = Field(description="Top 3-4 professional strengths highlighted in the resume that match the job description")
    gaps: List[str] = Field(description="Key gaps in experience, skills, certifications or qualifications required by the job description")
    recommendations: List[str] = Field(description="Actionable, step-by-step suggestions to modify the resume and improve the match score")
    keywords: List[KeywordMatch] = Field(description="List of key technical and soft skills/keywords found or missing")
    bullet_improvements: List[BulletPointImprovement] = Field(description="3-4 before/after examples of how to rewrite specific resume bullets to be more effective and JD-aligned")

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text content from PDF file bytes."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text.strip()
    except Exception as e:
        raise ValueError(f"Error reading PDF file: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text content from DOCX file bytes."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        text = []
        # Extract from paragraphs
        for para in doc.paragraphs:
            if para.text:
                text.append(para.text)
        # Extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text.append(cell.text)
        return "\n".join(text).strip()
    except Exception as e:
        raise ValueError(f"Error reading DOCX file: {str(e)}")

def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Detects file type and extracts text."""
    ext = os.path.splitext(filename.lower())[1]
    if ext == '.pdf':
        return extract_text_from_pdf(file_bytes)
    elif ext in ['.docx', '.doc']:
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")

def analyze_resume(resume_text: str, jd_text: str, api_key: Optional[str] = None) -> ResumeAnalysis:
    """
    Sends the resume text and job description to Gemini 2.5 Flash
    and returns a structured ResumeAnalysis object.
    """
    # Fallback to env var if key is not passed
    effective_api_key = api_key or os.environ.get("GEMINI_API_KEY")
    if not effective_api_key:
        raise ValueError("Gemini API key is required. Please set the GEMINI_API_KEY environment variable or pass it in the request.")

    # Initialize Google GenAI client
    client = genai.Client(api_key=effective_api_key)

    prompt = f"""
You are an expert ATS (Applicant Tracking System) parser and senior recruiter.
Analyze the candidate's resume below against the provided Job Description (JD).
Your evaluation must be objective, meticulous, and provide actionable, high-impact improvements.

CRITICAL SCORING LOGIC FOR QUALIFICATIONS AND EXPERIENCE:
You must FIRST extract and evaluate the following:
A. Mandatory qualifications
B. Mandatory certifications/licences
C. Required technical/domain skills
D. Required industry-specific experience
E. Required years/type of experience
F. Preferred qualifications
G. Transferable skills

Evaluate the candidate in this strict order:
1. Mandatory/core qualifications
2. Required technical or domain-specific skills
3. Required industry-specific experience
4. Required certifications/licences, if applicable
5. Relevant job-specific experience
6. Transferable skills
7. General/soft skills

- CORE MISMATCH: If a mandatory qualification, required domain skill, or required industry experience is clearly missing, this is a CRITICAL GAP. Do NOT allow transferable skills to compensate for a missing mandatory/core qualification. The match score MUST drop significantly (HIGH IMPACT).
- PREFERRED VS MANDATORY: Distinguish between REQUIRED/MANDATORY and PREFERRED qualifications. Missing a preferred qualification should have a MEDIUM/LOW IMPACT. Do NOT automatically reject or heavily penalize a candidate from a different industry if the JD says industry experience is "preferred" or not strictly required.
- EXPLANATION REQUIREMENT: When there is a major qualification mismatch, clearly explain in your summary or gaps:
  1. What the job requires.
  2. What the resume shows.
  3. What is missing or mismatched.
  4. Whether the requirement appears mandatory or preferred.
  5. How this affects the overall match.
- NO HALLUCINATION: ONLY use information found in the resume and JD. NEVER invent qualifications, experience, skills, etc. If something is missing, explicitly state "The qualification is not demonstrated in the provided resume." Do not say the candidate definitely does not have it unless the available evidence supports that conclusion.

Resume Text:
{resume_text}

Job Description:
{jd_text}

Provide:
1. An overall match score (0-100) and rating category.
2. Dimension scores (Skills Alignment, Experience Relevance, Formatting & Structure, Education) with specific feedback.
3. Top strengths matching the JD.
4. Core gaps where the resume fails to meet the JD criteria. If a mandatory qualification is missing, explicitly flag it as a critical gap and explain the mismatch.
5. Specific actionable recommendations for changes.
6. A list of critical keywords and skills from the JD (identify if they are present or missing, and their importance level).
7. Specific before-and-after bullet point rewrites. For weak bullets, identify why they are weak and provide a stronger rewritten version using strong action verbs. Make the wording more specific where the source resume provides enough info. NEVER invent numbers, percentages, achievements, responsibilities, companies, or experience. If there is not enough info to create a genuinely specific or quantified bullet, suggest what measurable/quantifiable info the user should add. Clearly distinguish between a safe rewrite based on existing info, and additional info the candidate could add themselves.
"""

    response = client.models.generate_content(
        model='gemini-2.5-flash',
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=ResumeAnalysis,
            temperature=0.1,  # Low temperature for analytical consistency
        ),
    )

    # The SDK returns the structured output parsed into the model specified in response_schema
    # Under the hood, response.parsed contains the schema-validated object
    if hasattr(response, 'parsed') and response.parsed:
        return response.parsed
    else:
        # Fallback in case parsed is missing
        import json
        data = json.loads(response.text)
        return ResumeAnalysis(**data)
